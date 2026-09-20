# -*- coding: utf-8 -*-
"""Reconstrução de Marcos de Desenvolvimento/Marcos.xlsx (padrão do Resumo Relatórios)."""
from __future__ import annotations

import calendar
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from caminhos import MARCOS  # noqa: E402

NASC = date(2020, 6, 30)
OUT = MARCOS / "Marcos.xlsx"

# Mesmas famílias de cor do Resumo Relatórios / Resumo Exames
CAT_COLORS = {
    "Saúde": "D6EAF8",
    "Família / religião": "FAE5D3",
    "Desenvolvimento motor": "E8DAEF",
    "Cuidados pessoais": "F5E6D3",
    "Comunicação": "FCF3CF",
    "Sono / rotina": "D4E6F1",
    "Alimentação": "D5F5E3",
    "Desenvolvimento social": "FADBD8",
    "Desenvolvimento físico": "F5CBA7",
    "Autonomia": "FDEBD0",
}

# (data, categoria, marco) — textos conferidos nas capturas de BackUp/
MARCOS_ROWS: list[tuple[date, str, str]] = [
    (date(2020, 8, 4), "Saúde", "Primeira cirurgia DVE"),
    (date(2020, 8, 12), "Família / religião", "Batismo"),
    (date(2020, 8, 13), "Família / religião", "Convite para padrinhos de batismo"),
    (date(2020, 8, 18), "Saúde", "Segunda cirurgia DVE"),
    (date(2020, 9, 8), "Saúde", "Terceira cirurgia DVP"),
    (date(2020, 10, 1), "Saúde", "Alta hospitalar"),
    (date(2020, 11, 19), "Desenvolvimento motor", "Desvirou sozinha"),
    (date(2020, 12, 10), "Cuidados pessoais", "Furou a orelha"),
    (date(2020, 12, 18), "Família / religião", "Convite de madrinha para consagração"),
    (date(2020, 12, 19), "Família / religião", "Convite de padrinho para consagração"),
    (date(2020, 12, 24), "Saúde", "Última dose de Gardenal"),
    (date(2020, 12, 27), "Desenvolvimento motor", "Virou sozinha"),
    (date(2020, 12, 31), "Comunicação", "Chamou pelo colo com os braços"),
    (date(2021, 1, 2), "Desenvolvimento motor", "Tirou a fralda do rosto"),
    (date(2021, 1, 15), "Sono / rotina", "Dormiu a noite sem o ninho"),
    (date(2021, 1, 16), "Alimentação", "Introdução alimentar"),
    (date(2021, 1, 17), "Desenvolvimento social", "Sorriso simpático"),
    (date(2021, 1, 18), "Desenvolvimento motor", "Bateu palminhas para a bisa"),
    (date(2021, 1, 22), "Desenvolvimento motor", "Rolou pela primeira vez"),
    (date(2021, 2, 2), "Desenvolvimento motor", "Sentou sem apoio pela primeira vez"),
    (date(2021, 2, 19), "Desenvolvimento social", "Beijou pela primeira vez"),
    (date(2021, 3, 8), "Comunicação", "Deu tchauzinho pela primeira vez"),
    (date(2021, 3, 8), "Sono / rotina", "Acordou de madrugada, pesadelo"),
    (date(2021, 3, 12), "Desenvolvimento motor", "Arrastou-se pela primeira vez"),
    (date(2021, 3, 29), "Desenvolvimento motor", "Ficou em posição de esfinge"),
    (date(2021, 3, 30), "Desenvolvimento motor", "Escondeu sozinha o rosto"),
    (date(2021, 4, 7), "Comunicação", "Falou “mamãe”"),
    (date(2021, 4, 21), "Desenvolvimento motor", "Arrastou-se"),
    (date(2021, 5, 10), "Alimentação", "Comeu tudo e raspou o prato"),
    (date(2021, 5, 21), "Desenvolvimento motor", "Sentada, remou para trás"),
    (date(2021, 5, 23), "Desenvolvimento motor", "Saiu da posição sentada para deitada"),
    (date(2021, 5, 26), "Comunicação", "Chamou a água de “Iá”. Várias vezes"),
    (date(2021, 5, 27), "Alimentação", "Comeu com interesse o frango"),
    (date(2021, 7, 19), "Desenvolvimento motor", "Saiu da posição deitada para sentada"),
    (date(2021, 8, 11), "Desenvolvimento motor", "Ficou em pé com apoio"),
    (date(2021, 8, 17), "Desenvolvimento motor", "Engatinhou no banheiro pela primeira vez"),
    (date(2021, 8, 19), "Desenvolvimento físico", "Dentinho de baixo apontou"),
    (date(2021, 11, 13), "Desenvolvimento motor", "Colocou-se em pé sozinha pela primeira vez"),
    (date(2022, 2, 9), "Desenvolvimento motor", "Andou sem apoio 1 vez (depois quis colo)"),
    (date(2024, 2, 10), "Autonomia", "Desfralde"),
]


def _aniversario_mensal(nasc: date, year: int, month: int) -> date:
    """Dia do aniversário mensal (30/06 → dia 30, ou último dia do mês se não existir)."""
    return date(year, month, min(nasc.day, calendar.monthrange(year, month)[1]))


def idade_ymd(nasc: date, d: date) -> tuple[int, int, int]:
    """Idade no mesmo critério do diário Momentos (aniversário mensal + dias restantes)."""
    ann = _aniversario_mensal(nasc, d.year, d.month)
    if ann > d:
        if d.month == 1:
            ann = _aniversario_mensal(nasc, d.year - 1, 12)
        else:
            ann = _aniversario_mensal(nasc, d.year, d.month - 1)
    months = (ann.year - nasc.year) * 12 + (ann.month - nasc.month)
    days = (d - ann).days
    return months // 12, months % 12, days


def _parte(n: int, singular: str, plural: str) -> str | None:
    if n <= 0:
        return None
    return f"{n} {singular if n == 1 else plural}"


def idade_pt(d: date) -> str:
    y, m, day = idade_ymd(NASC, d)
    parts = [
        p
        for p in (
            _parte(y, "ano", "anos"),
            _parte(m, "mês", "meses"),
            _parte(day, "dia", "dias"),
        )
        if p
    ]
    if not parts:
        return "nascimento"
    if len(parts) == 1:
        return parts[0]
    if len(parts) == 2:
        return f"{parts[0]} e {parts[1]}"
    return f"{parts[0]}, {parts[1]} e {parts[2]}"


def semana_vida(d: date) -> int:
    return (d - NASC).days // 7 + 1


def _style_header_row(ws, row: int, ncols: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    for col in range(1, ncols + 1):
        cell = ws.cell(row, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin
    ws.row_dimensions[row].height = 20


def _print_setup(ws) -> None:
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)


def build() -> Path:
    MARCOS.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Marcos"
    ws.sheet_view.showGridLines = False

    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    body_font = Font(name="Calibri", size=10, color="333333")
    no_fill = PatternFill(fill_type=None)

    ws.merge_cells("A1:E1")
    title = ws["A1"]
    title.value = "Cecília Maria Albergaria Silva — Marcos de desenvolvimento"
    title.font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    title.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells("A2:E2")
    sub = ws["A2"]
    sub.value = (
        f"Nasc. {NASC.strftime('%d/%m/%Y')}  ·  "
        f"{len(MARCOS_ROWS)} registros do diário Momentos  ·  "
        "ago/2020 a fev/2024"
    )
    sub.font = Font(name="Calibri", size=10, italic=True, color="666666")
    sub.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 18

    headers = ["Data", "Idade", "Semana", "Categoria", "Marco"]
    for col, h in enumerate(headers, 1):
        ws.cell(3, col, h)
    _style_header_row(ws, 3, 5)

    rows = sorted(MARCOS_ROWS, key=lambda x: (x[0], x[2]))
    for i, (d, cat, marco) in enumerate(rows):
        r = i + 4
        ws.cell(r, 1, d).number_format = "DD/MM/YYYY"
        ws.cell(r, 2, idade_pt(d))
        ws.cell(r, 3, semana_vida(d))
        ws.cell(r, 4, cat)
        ws.cell(r, 5, marco)

        tipo_fill = PatternFill("solid", fgColor=CAT_COLORS.get(cat, "FFFFFF"))
        for col in range(1, 6):
            cell = ws.cell(r, col)
            cell.font = body_font
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 5))
            if col == 4:
                cell.fill = tipo_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.fill = no_fill
                if col in (1, 3):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = min(48, max(18, 16 + (len(marco) // 55) * 12))

    last = 3 + len(rows)

    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 24
    ws.column_dimensions["E"].width = 58
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:E{last}"
    _print_setup(ws)

    note = ws.cell(
        last + 2,
        1,
        "Fonte: diário Momentos (capturas em BackUp/). "
        "Legendas cortadas no app foram completadas pelas fotos. "
        "DVE = derivação ventricular externa; DVP = derivação ventrículo-peritoneal.",
    )
    note.font = Font(name="Calibri", size=9, italic=True, color="666666")
    ws.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=5)
    ws.row_dimensions[last + 2].height = 28

    # --- aba Por categoria ---
    ws2 = wb.create_sheet("Por categoria")
    ws2.sheet_view.showGridLines = False
    ws2.merge_cells("A1:D1")
    t2 = ws2["A1"]
    t2.value = "Cecília Maria Albergaria Silva — Marcos por categoria"
    t2.font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    t2.alignment = Alignment(vertical="center")
    ws2.row_dimensions[1].height = 26

    ws2.merge_cells("A2:D2")
    s2 = ws2["A2"]
    s2.value = f"Nasc. {NASC.strftime('%d/%m/%Y')}  ·  {len(rows)} marcos no total"
    s2.font = Font(name="Calibri", size=10, italic=True, color="666666")
    ws2.row_dimensions[2].height = 18

    for col, h in enumerate(["Categoria", "Quantidade", "Primeiro", "Último"], 1):
        ws2.cell(3, col, h)
    _style_header_row(ws2, 3, 4)

    first_last: dict[str, tuple[date, date]] = {}
    counts = Counter()
    for d, cat, _marco in rows:
        counts[cat] += 1
        if cat not in first_last:
            first_last[cat] = (d, d)
        else:
            a, b = first_last[cat]
            first_last[cat] = (min(a, d), max(b, d))

    ordem = sorted(counts, key=lambda c: (-counts[c], c))
    for i, cat in enumerate(ordem):
        r = i + 4
        ws2.cell(r, 1, cat)
        ws2.cell(r, 2, counts[cat])
        ws2.cell(r, 3, first_last[cat][0]).number_format = "DD/MM/YYYY"
        ws2.cell(r, 4, first_last[cat][1]).number_format = "DD/MM/YYYY"
        tipo_fill = PatternFill("solid", fgColor=CAT_COLORS.get(cat, "FFFFFF"))
        for col in range(1, 5):
            cell = ws2.cell(r, col)
            cell.font = body_font
            cell.border = thin
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if col == 1:
                cell.fill = tipo_fill
            else:
                cell.fill = no_fill
        ws2.row_dimensions[r].height = 20

    last2 = 3 + len(ordem)
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 14
    ws2.column_dimensions["D"].width = 14
    ws2.freeze_panes = "A4"
    ws2.auto_filter.ref = f"A3:D{last2}"
    _print_setup(ws2)

    if OUT.exists():
        try:
            OUT.unlink()
        except OSError:
            # Excel pode estar com o arquivo aberto — grava ao lado
            alt = MARCOS / "Marcos - novo.xlsx"
            wb.save(alt)
            wb.close()
            print("LOCKED", OUT)
            print("SAVED", alt)
            return alt

    wb.save(OUT)
    wb.close()
    print("SAVED", OUT)
    print("ROWS", len(rows))
    return OUT


if __name__ == "__main__":
    build()
