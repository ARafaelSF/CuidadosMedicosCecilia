# -*- coding: utf-8 -*-
"""Comparativo idade óssea: laudo × IA → planilha com gráfico + PDF para imprimir."""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
EXAMES = ROOT / "Exames"
XLSX = EXAMES / "Idade Ossea.xlsx"
PDF_OUT = EXAMES / "Idade Ossea - Comparativo laudo vs IA.pdf"
DOB = date(2020, 6, 30)


def parse_ym(s: str | None) -> float | None:
    if not s or s == "—":
        return None
    m = re.match(r"(\d+)a(\d+)m", str(s).strip())
    if not m:
        return None
    return int(m.group(1)) + int(m.group(2)) / 12


def load_rows() -> list[dict]:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Idade óssea"]
    rows: list[dict] = []
    for r in range(5, ws.max_row + 1):
        d = ws.cell(r, 1).value
        crono_txt = ws.cell(r, 2).value
        crono_anos = ws.cell(r, 5).value
        if not isinstance(d, datetime) or not crono_txt or crono_anos is None:
            continue
        exam = d.date()
        rows.append({
            "data": exam,
            "crono_txt": ws.cell(r, 2).value,
            "ossea_txt": ws.cell(r, 3).value,
            "atraso_laudo": ws.cell(r, 4).value,
            "crono_anos": float(crono_anos),
            "ossea_anos": float(ws.cell(r, 6).value),
            "ia_txt": ws.cell(r, 7).value,
            "atraso_ia": ws.cell(r, 8).value,
            "ia_anos": ws.cell(r, 9).value if ws.cell(r, 9).value not in (None, "—") else None,
            "concordancia": ws.cell(r, 10).value,
            "notas": ws.cell(r, 11).value,
        })
    wb.close()
    return rows


def ensure_grafico_sheet(rows: list[dict]) -> None:
    wb = load_workbook(XLSX)
    if "Gráfico" in wb.sheetnames:
        del wb["Gráfico"]
    ws = wb.create_sheet("Gráfico", 1)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )

    ws.merge_cells("A1:H1")
    ws["A1"] = "Cecília — Idade óssea: laudo × IA (Greulich-Pyle, feminino)"
    ws["A1"].font = title_font
    ws.row_dimensions[1].height = 24

    headers = [
        "Data", "Cronológica", "Óssea laudo", "Óssea IA", "Atraso laudo (m)",
        "Atraso IA (m)", "Concordância", "Notas",
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    ws.row_dimensions[3].height = 28

    body = Font(name="Calibri", size=10)
    for i, row in enumerate(rows):
        r = 4 + i
        ia = row["ia_anos"]
        ws.cell(r, 1, row["data"]).number_format = "DD/MM/YYYY"
        ws.cell(r, 2, row["crono_txt"])
        ws.cell(r, 3, row["ossea_txt"])
        ws.cell(r, 4, row["ia_txt"] if ia is not None else "—")
        ws.cell(r, 5, row["atraso_laudo"])
        ws.cell(r, 6, row["atraso_ia"] if ia is not None else "—")
        ws.cell(r, 7, row["concordancia"])
        ws.cell(r, 8, row["notas"])
        for c in range(1, 9):
            cell = ws.cell(r, c)
            cell.font = body
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=(c >= 7))

    # dados numéricos para gráfico (anos decimais)
    start = 4 + len(rows) + 2
    ws.cell(start, 1, "Data (gráfico)")
    ws.cell(start, 2, "Cronológica (anos)")
    ws.cell(start, 3, "Óssea laudo (anos)")
    ws.cell(start, 4, "Óssea IA (anos)")
    for c in range(1, 5):
        ws.cell(start, c).font = Font(bold=True, size=10)

    first = start + 1
    for i, row in enumerate(rows):
        r = first + i
        ws.cell(r, 1, datetime.combine(row["data"], datetime.min.time()))
        ws.cell(r, 1).number_format = "DD/MM/YYYY"
        ws.cell(r, 2, row["crono_anos"])
        ws.cell(r, 3, row["ossea_anos"])
        ia = row["ia_anos"]
        if ia is not None:
            ws.cell(r, 4, float(ia))

    last = first + len(rows) - 1
    chart = LineChart()
    chart.title = "Idade cronológica × idade óssea (laudo e IA)"
    chart.style = 2
    chart.y_axis.title = "Anos"
    chart.x_axis.title = "Data do exame"
    chart.height = 12
    chart.width = 22
    chart.legend.position = "b"

    cats = Reference(ws, min_col=1, min_row=first, max_row=last)
    for col in (2, 3, 4):
        vals = Reference(ws, min_col=col, min_row=start, max_row=last)
        chart.add_data(vals, titles_from_data=True)
    chart.set_categories(cats)

    ws.add_chart(chart, f"A{last + 3}")

    for col, w in {"A": 12, "B": 14, "C": 14, "D": 14, "E": 14, "F": 12, "G": 16, "H": 52}.items():
        ws.column_dimensions[col].width = w

    wb.save(XLSX)
    wb.close()


def _conc_short(s: str | None) -> str:
    if not s:
        return "—"
    m = {
        "Concordante": "Concordante",
        "Parcialmente divergente": "Parc. divergente",
        "Sem RX no PDF": "Sem RX",
    }
    return m.get(s, s)


def make_pdf(rows: list[dict]) -> None:
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.size": 10,
        "axes.titlesize": 11,
        "axes.labelsize": 10,
    })

    dates = [datetime.combine(r["data"], datetime.min.time()) for r in rows]
    crono = [r["crono_anos"] for r in rows]
    laudo = [r["ossea_anos"] for r in rows]
    ia = [float(r["ia_anos"]) if r["ia_anos"] is not None else None for r in rows]

    fig = plt.figure(figsize=(11.69, 8.27))
    fig.patch.set_facecolor("white")

    # Cabeçalho (fora da área do gráfico)
    fig.text(
        0.5, 0.965,
        "Cecília Maria Albergaria Silva — Idade óssea (RX punho)",
        ha="center", va="top", fontsize=14, fontweight="bold", color="#1F4E79",
    )
    fig.text(
        0.5, 0.935,
        f"Nascimento: {DOB.strftime('%d/%m/%Y')}  ·  Greulich-Pyle (feminino)",
        ha="center", va="top", fontsize=9, color="#555",
    )
    fig.text(
        0.5, 0.915,
        "Laudo radiologista vs estimativa visual (IA)",
        ha="center", va="top", fontsize=9, color="#555",
    )

    # Gráfico — começa abaixo do cabeçalho
    ax = fig.add_axes([0.08, 0.43, 0.88, 0.43])
    ax.plot(dates, crono, "o-", color="#1F4E79", linewidth=2, markersize=8, label="Idade cronológica")
    ax.plot(dates, laudo, "D-", color="#C0392B", linewidth=2, markersize=8, label="Idade óssea (laudo)")
    ia_dates = [d for d, v in zip(dates, ia) if v is not None]
    ia_vals = [v for v in ia if v is not None]
    if ia_dates:
        ax.plot(ia_dates, ia_vals, "^--", color="#27AE60", linewidth=2, markersize=9, label="Idade óssea (IA)")
    ax.set_ylabel("Anos")
    ax.set_xlabel("Data do exame", labelpad=14)
    ax.set_title("Evolução: cronológica × óssea (laudo e IA)", pad=12, fontsize=11)
    ax.legend(loc="upper left", framealpha=0.95, fontsize=9)
    ax.grid(True, alpha=0.3)
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%m/%Y"))
    ax.xaxis.set_major_locator(mdates.AutoDateLocator())
    ymax = max(crono + laudo + [v for v in ia if v is not None]) + 0.8
    ax.set_ylim(bottom=0, top=ymax)
    for label in ax.get_xticklabels():
        label.set_rotation(30)
        label.set_ha("right")

    # Tabela resumo
    ax_tbl = fig.add_axes([0.06, 0.14, 0.88, 0.17])
    ax_tbl.axis("off")
    col_labels = ["Data", "Cron.", "Óssea\nlaudo", "Óssea\nIA", "Atraso\nlaudo", "Atraso\nIA", "Concord."]
    table_data = []
    for r in rows:
        table_data.append([
            r["data"].strftime("%d/%m/%Y"),
            r["crono_txt"],
            r["ossea_txt"],
            r["ia_txt"] if r["ia_anos"] is not None else "—",
            f"{r['atraso_laudo']} m",
            f"{r['atraso_ia']} m" if r["ia_anos"] is not None else "—",
            _conc_short(r["concordancia"]),
        ])
    col_widths = [0.13, 0.09, 0.11, 0.11, 0.11, 0.11, 0.14]
    tbl = ax_tbl.table(
        cellText=table_data,
        colLabels=col_labels,
        loc="center",
        cellLoc="center",
        colWidths=col_widths,
    )
    tbl.auto_set_font_size(False)
    tbl.set_fontsize(8.5)
    tbl.scale(1, 1.45)
    for (row, col), cell in tbl.get_celld().items():
        cell.set_height(0.18)
        if row == 0:
            cell.set_facecolor("#1F4E79")
            cell.set_text_props(color="white", fontweight="bold", fontsize=8)
        elif row % 2 == 0:
            cell.set_facecolor("#F4F6F7")
        if col == 6:
            cell.get_text().set_fontsize(7.5)

    fig.text(
        0.06, 0.06,
        "IA = leitura visual assistida; não substitui laudo.\n"
        "Fev/2025: PDF sem radiografia.  Jul/2026: IA ~5a2m vs laudo 4a2m (epífise ulnar visível no RX).",
        ha="left", va="bottom", fontsize=8, color="#666", linespacing=1.35,
    )

    with PdfPages(PDF_OUT) as pdf:
        pdf.savefig(fig, dpi=150)
    plt.close(fig)


def main() -> None:
    import argparse

    ap = argparse.ArgumentParser()
    ap.add_argument("--so-pdf", action="store_true", help="Regenera só o PDF (não altera a planilha).")
    args = ap.parse_args()

    rows = load_rows()
    if not rows:
        raise SystemExit("Nenhuma linha em Idade Ossea.xlsx")
    if not args.so_pdf:
        try:
            ensure_grafico_sheet(rows)
            print(f"OK planilha: {XLSX} (aba Gráfico)")
        except PermissionError:
            print(f"AVISO: planilha aberta — PDF será gerado; feche o Excel e rode de novo para atualizar a aba Gráfico.")
    make_pdf(rows)
    print(f"OK imprimir: {PDF_OUT}")


if __name__ == "__main__":
    main()
